#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.13"
# dependencies = [
#   "python-docx>=1.2.0,<2",
#   "reportlab>=4.4.0,<5",
# ]
# ///
"""Create synthetic PDF and DOCX files for the document review demo."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIRECTORY = ROOT / "samples" / "documents"
TITLE = "Project Atlas Launch Review"
SECTIONS = [
    (
        "Summary",
        "Project Atlas is a customer onboarding workflow planned for a limited launch.",
    ),
    (
        "Decision",
        "The review board approved Project Atlas for a limited launch on October 15.",
    ),
    (
        "Key change",
        "The launch scope now includes North America only; Europe moves to a later phase.",
    ),
    (
        "Risk",
        "The vendor support agreement does not yet define weekend response targets.",
    ),
    (
        "Open question",
        "The document does not identify who will approve expansion beyond North America.",
    ),
    (
        "Review focus",
        "Reviewers should confirm the weekend support target before the launch checklist is finalized.",
    ),
]


def _create_docx(path: Path) -> None:
    document = Document()
    document.add_heading(TITLE, level=0)
    document.add_paragraph("Synthetic sample - contains no customer data.")
    for heading, text in SECTIONS:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)
    document.save(path)


def _create_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    story = [
        Paragraph(TITLE, styles["Title"]),
        Paragraph("Synthetic sample - contains no customer data.", styles["Italic"]),
        Spacer(1, 12),
    ]
    for heading, text in SECTIONS:
        story.extend(
            (
                Paragraph(heading, styles["Heading2"]),
                Paragraph(text, styles["BodyText"]),
                Spacer(1, 8),
            )
        )
    document = SimpleDocTemplate(str(path), pagesize=LETTER)
    document.build(story)


def main() -> int:
    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    docx_path = OUTPUT_DIRECTORY / "project-atlas-review.docx"
    pdf_path = OUTPUT_DIRECTORY / "project-atlas-review.pdf"
    _create_docx(docx_path)
    _create_pdf(pdf_path)
    print(f"Created {docx_path.relative_to(ROOT)}")
    print(f"Created {pdf_path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
